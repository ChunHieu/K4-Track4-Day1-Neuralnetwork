# -*- coding: utf-8 -*-
"""results_table.py — Hoàn thiện từ pseudo-code.

Nhiệm vụ: lưu kết quả từng lần chạy ra JSON, rồi điền vào experiments.xlsx từ mẫu
templates/experiment_table_template.xlsx (đừng gõ tay hàng chục dòng, rất dễ sai).

Tên cột của sheet "Experiments" (giữ nguyên, đúng thứ tự mẫu):
    exp_id, group, description, loss, optimizer, lr, weight_decay, batch, epochs, hidden, dropout,
    clip_norm, precision, init, seed, step0_loss, best_val_loss, best_epoch, final_train_loss,
    final_val_loss, val_acc, val_macro_f1, time_per_epoch_s, peak_mem_MB, diverged,
    eval_acc, eval_macro_f1, figure_file, notes
(các cột công thức ở cuối bảng mẫu tự tính, đừng ghi đè)
"""
from __future__ import annotations

import json
from pathlib import Path

# Danh sách cột dữ liệu (không phải cột công thức) theo đúng thứ tự trong template
_DATA_COLS = [
    "exp_id", "group", "description", "loss", "optimizer", "lr", "weight_decay",
    "batch", "epochs", "hidden", "dropout", "clip_norm", "precision", "init", "seed",
    "step0_loss", "best_val_loss", "best_epoch", "final_train_loss", "final_val_loss",
    "val_acc", "val_macro_f1", "time_per_epoch_s", "peak_mem_MB", "diverged",
    "eval_acc", "eval_macro_f1", "figure_file", "notes",
]

# Cột công thức trong template — KHÔNG ghi đè
_FORMULA_COLS = {"step0_gap_vs_lnC", "gap_val_minus_train", "delta_val_f1_vs_base", "beyond_noise"}


def save_result(result: dict, results_dir: str = "../results") -> str:
    """Ghi result["cfg"], result["history"], result["summary"] (KHÔNG ghi best_state) ra
    <results_dir>/<exp_id>.json. Trả về đường dẫn file. Tạo thư mục nếu chưa có."""
    out_dir = Path(results_dir)
    out_dir.mkdir(parents=True, exist_ok=True)

    exp_id = result["cfg"]["exp_id"]
    payload = {
        "cfg":     result["cfg"],
        "history": result["history"],
        "summary": result["summary"],
    }

    # Chuyển các kiểu không JSON-serializable
    def _serialize(obj):
        if isinstance(obj, (bool,)):
            return obj
        if hasattr(obj, "item"):            # numpy scalar
            return obj.item()
        if hasattr(obj, "tolist"):          # numpy array / torch tensor
            return obj.tolist()
        raise TypeError(f"Không thể serialize {type(obj)}: {obj!r}")

    out_path = out_dir / f"{exp_id}.json"
    with open(out_path, "w", encoding="utf-8") as f:
        json.dump(payload, f, ensure_ascii=False, indent=2, default=_serialize)

    print(f"  Saved result -> {out_path}")
    return str(out_path)


def load_results(results_dir: str = "../results") -> list[dict]:
    """Đọc mọi file *.json trong results_dir, trả về danh sách dict (sắp theo exp_id)."""
    out_dir = Path(results_dir)
    if not out_dir.exists():
        return []
    results = []
    for fp in sorted(out_dir.glob("*.json")):
        with open(fp, "r", encoding="utf-8") as f:
            results.append(json.load(f))
    # Sắp xếp theo exp_id
    results.sort(key=lambda r: r.get("cfg", {}).get("exp_id", ""))
    return results


def to_row(result: dict, eval_scores: dict | None = None, notes: str = "") -> dict:
    """Biến một kết quả thành một dòng của bảng: gộp cfg + summary (+ eval_acc, eval_macro_f1 nếu có)
    + figure_file = f"figures/{exp_id}.png". Khoá phải trùng tên cột ở đầu file.
    Chỉ truyền eval_scores cho baseline và cấu hình cuối cùng.
    """
    cfg     = result.get("cfg", {})
    summary = result.get("summary", {})
    exp_id  = cfg.get("exp_id", "")

    row = {}
    # Từ cfg
    for key in ["exp_id", "group", "description", "loss", "optimizer", "lr",
                "weight_decay", "batch", "epochs", "hidden", "dropout",
                "clip_norm", "precision", "init", "seed"]:
        val = cfg.get(key)
        # hidden là tuple → chuyển sang string "256-128"
        if key == "hidden" and val is not None:
            val = "-".join(str(h) for h in val)
        row[key] = val

    # Từ summary
    for key in ["step0_loss", "best_val_loss", "best_epoch", "final_train_loss",
                "final_val_loss", "val_acc", "val_macro_f1",
                "time_per_epoch_s", "peak_mem_MB", "diverged"]:
        row[key] = summary.get(key)

    # eval_acc / eval_macro_f1 (chỉ baseline và cấu hình cuối)
    if eval_scores is not None:
        row["eval_acc"]      = eval_scores.get("accuracy") or eval_scores.get("eval_acc")
        row["eval_macro_f1"] = eval_scores.get("macro_f1") or eval_scores.get("eval_macro_f1")
    else:
        row["eval_acc"]      = None
        row["eval_macro_f1"] = None

    row["figure_file"] = f"figures/{exp_id}.png"
    row["notes"]       = notes

    return row


def write_xlsx(rows: list[dict], template_path: str, out_path: str) -> None:
    """Điền các dòng vào sheet "Experiments" của mẫu, từ dòng 2 trở xuống, rồi lưu thành out_path.

    Các bước (openpyxl):
      1. wb = openpyxl.load_workbook(template_path)   # KHÔNG dùng data_only=True
      2. ws = wb["Experiments"]; đọc tiêu đề dòng 1
      3. với mỗi row: ghi giá trị vào đúng cột; BỎ QUA cột công thức
      4. wb.save(out_path)
    """
    import openpyxl  # type: ignore

    wb = openpyxl.load_workbook(template_path)
    ws = wb["Experiments"]

    # Đọc tiêu đề dòng 1 → ánh xạ tên cột -> số cột (1-indexed)
    header_map: dict[str, int] = {}
    for col_idx, cell in enumerate(ws[1], start=1):
        if cell.value is not None:
            header_map[str(cell.value).strip()] = col_idx

    # Ghi dữ liệu từ dòng 2
    for row_idx, row_data in enumerate(rows, start=2):
        for col_name, col_idx in header_map.items():
            # Bỏ qua cột công thức
            if col_name in _FORMULA_COLS:
                continue
            val = row_data.get(col_name)
            # Chuyển None → chuỗi rỗng, bool giữ nguyên, còn lại giữ nguyên
            if val is None:
                val = ""
            ws.cell(row=row_idx, column=col_idx, value=val)

    out = Path(out_path)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    print(f"  Saved table -> {out_path} ({len(rows)} rows)")
